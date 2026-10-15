"""Cover letters: an honesty check on the text, and PDF/DOCX output that passes the
same ATS checks as the CV.

The client's LLM drafts the letter from the person's bullets and skills. This
module only checks it and renders it, deterministically.

`check_claims` has two kinds of findings:
- problems block the build: the letter names a skill the inventory marks X, or a
  familiarity (F) skill without saying it is only familiarity;
- warnings are for the person to review: a technology the inventory and the
  bullet pool don't show, or a number that appears nowhere in the profile.
"""

from __future__ import annotations

import io
import json
import re
from dataclasses import asdict, dataclass, field
from datetime import date
from pathlib import Path

import typst
from docx import Document as DocxDocument
from docx.shared import Inches, Pt

from honest_cv import checklist
from honest_cv.document import contact_items
from honest_cv.model import Lang, Level, Source
from honest_cv.render.pdf import FONTS
from honest_cv.score import _match, _pattern, find_skill, groups, spellings
from honest_cv.verify import Check, Report, verify_pdf

TEMPLATE = Path(__file__).parent / "render" / "templates" / "letter.typ"
FORMATS = ("pdf", "docx")

# Words that mark a mention as familiarity, not experience, in the same sentence.
FAMILIARITY_WORDS = re.compile(
    r"\b(familiar\w*|basic|exposure|learning|studied|explored|tried|introductory|"
    r"básic[oa]s?|nociones|aprendiendo|estudi[eé]|explor[eé]|prob[eé]|introductori[oa])\b",
    re.I,
)
SENTENCE = re.compile(r"(?<=[.!?;])\s+|\n+")
NUMBER = re.compile(r"(?<![\w.])\d+(?:[.,]\d+)*\s?(?:%|x|k|K|M)?(?![\w])")
# Checklists of practices (people, process, design), not technologies: in a letter,
# words like "scalability" or "mentoring" are not claims to check.
NOT_TECHNOLOGY = {"engineering-lead", "software-architecture"}

GREETING: dict[Lang, str] = {"en": "Dear Hiring Team,", "es": "Estimado equipo de selección:"}
CLOSING: dict[Lang, str] = {"en": "Sincerely,", "es": "Atentamente,"}
MONTHS: dict[Lang, list[str]] = {
    "en": ["January", "February", "March", "April", "May", "June", "July", "August",
           "September", "October", "November", "December"],
    "es": ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
           "septiembre", "octubre", "noviembre", "diciembre"],
}


@dataclass
class Finding:
    kind: str
    term: str
    sentence: str
    detail: str


@dataclass
class Claims:
    problems: list[Finding] = field(default_factory=list)
    warnings: list[Finding] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.problems

    def to_dict(self) -> dict:
        return {"ok": self.ok, "problems": [asdict(f) for f in self.problems],
                "warnings": [asdict(f) for f in self.warnings]}


def _vocabulary(source: Source, keywords: list[str]) -> list[str]:
    """Names to look for: synonyms, the technical built-in checklists, the person's own
    templates (any profession), the inventory and the post."""
    names = [g[0] for g in groups(source)]
    for b in checklist.available(source.folder):
        if b.id not in NOT_TECHNOLOGY or b.origin == "yours":
            names += [i.name for i in checklist.parse(Path(b.path).read_text("utf-8")).items]
    names += [s.name for s in source.inventory.skills] + keywords
    seen, out = set(), []
    for name in names:
        key = checklist._key(name, source)
        if key not in seen:
            seen.add(key)
            out.append(name)
    return out


def _bullet_text(source: Source) -> str:
    parts = []
    for b in source.bullets:
        parts += [b.text] if isinstance(b.text, str) else list(b.text.values())
        parts += b.tags + b.metrics
    return "\n".join(parts)


def _profile_numbers(source: Source) -> set[str]:
    text = [_bullet_text(source)]
    for r in source.profile.roles:
        text += [r.start, r.end or ""]
        text += [x for x in (r.context, r.title) if isinstance(x, str)]
        text += [v for x in (r.context, r.title) if isinstance(x, dict) for v in x.values()]
    for v in source.variants.values():
        for x in (v.summary, v.title):
            text += [x] if isinstance(x, str) else list((x or {}).values())
    for s in source.inventory.skills:
        text += [str(s.years or ""), str(s.last_used_year or ""), s.evidence or ""]
    return {_norm_number(n) for n in NUMBER.findall("\n".join(text))}


def _norm_number(n: str) -> str:
    return re.sub(r"[\s,]", "", n).replace("K", "k")


def _sentence(text: str, start: int) -> str:
    bounds = [0] + [m.end() for m in SENTENCE.finditer(text)] + [len(text)]
    lo = max(b for b in bounds if b <= start)
    hi = min((b for b in bounds if b > start), default=len(text))
    return text[lo:hi].strip()


def _inside(span: tuple[int, int], spans: list[tuple[int, int]]) -> bool:
    lo, hi = span
    return any(s <= lo and hi <= e and e - s > hi - lo for s, e in spans)


def check_claims(text: str, source: Source, keywords: list[str] | None = None) -> Claims:
    out = Claims()
    pool = _bullet_text(source)
    hits = {}
    for term in _vocabulary(source, keywords or []):
        names = spellings(term, source)
        if found := [m for name in names for m in _pattern(name).finditer(text)]:
            hits[term] = (names, found)
    spans = [m.span() for _, found in hits.values() for m in found]
    for term, (names, found) in hits.items():
        # "Spring" inside "Spring Boot" is a mention of Spring Boot, not of Spring.
        found = sorted((m for m in found if not _inside(m.span(), spans)), key=lambda m: m.start())
        if not found:
            continue
        sentence = _sentence(text, found[0].start())
        skill = find_skill(term, source)
        label = skill.name if skill else term
        if skill is not None and skill.level is Level.NONE:
            out.problems.append(Finding("not_held", label, sentence,
                                        "the inventory says the person does not have it (X)"))
        elif skill is not None and skill.level is Level.FAMILIARITY:
            # Every mention, not just the first: each sentence must say it is familiarity.
            for s in dict.fromkeys(_sentence(text, m.start()) for m in found):
                if not FAMILIARITY_WORDS.search(s):
                    out.problems.append(Finding("familiarity_as_experience", label, s,
                                                "it is familiarity (F): say so in this sentence, or leave it out"))
        elif skill is None and _match(names, pool) is None:
            out.warnings.append(Finding("not_in_inventory", label, sentence,
                                        "neither the inventory nor the bullets show it; ask the person"))
    known = _profile_numbers(source)
    for m in NUMBER.finditer(text):
        n = _norm_number(m.group())
        if len(n.rstrip("%xkM")) >= 2 and n not in known:
            out.warnings.append(Finding("number_not_in_profile", m.group().strip(), _sentence(text, m.start()),
                                        "this number appears nowhere in the profile: fine if it is about the company, "
                                        "otherwise check where it comes from"))
    return out


# --- rendering ----------------------------------------------------------------


@dataclass
class Letter:
    name: str
    contact: list[str]
    lang: Lang
    date: str
    company: str
    role: str
    recipient: str | None
    greeting: str
    paragraphs: list[str]
    closing: str

    def to_dict(self) -> dict:
        return asdict(self)


def today(lang: Lang, when: date | None = None) -> str:
    when = when or date.today()
    month = MONTHS[lang][when.month - 1]  # not strftime: it follows the system locale
    return f"{when.day} de {month} de {when.year}" if lang == "es" else f"{month} {when.day}, {when.year}"


def make(source: Source, lang: Lang, company: str, role: str, paragraphs: list[str], *,
         recipient: str | None = None, greeting: str | None = None, closing: str | None = None,
         when: date | None = None) -> Letter:
    body = [p.strip() for p in paragraphs if p and p.strip()]
    if not body:
        raise ValueError("the letter has no paragraphs")
    return Letter(
        name=source.profile.contact.name, contact=contact_items(source, lang), lang=lang,
        date=today(lang, when), company=company.strip(), role=role.strip(),
        recipient=recipient.strip() if recipient else None, greeting=greeting or GREETING[lang],
        paragraphs=body, closing=closing or CLOSING[lang],
    )


def render_pdf(letter: Letter) -> bytes:
    return typst.compile(
        str(TEMPLATE), font_paths=[str(FONTS)], ignore_system_fonts=True,
        sys_inputs={"data": json.dumps(letter.to_dict(), ensure_ascii=False)},
    )


def render_docx(letter: Letter) -> bytes:
    out = DocxDocument()
    for s in out.sections:
        s.top_margin = s.bottom_margin = Inches(0.8)
        s.left_margin = s.right_margin = Inches(0.9)
    normal = out.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    normal.paragraph_format.space_after = Pt(8)

    def para(text: str, bold: bool = False, size: float | None = None) -> None:
        run = out.add_paragraph().add_run(text)
        run.bold = bold
        if size:
            run.font.size = Pt(size)

    para(letter.name, bold=True, size=16)
    para(" · ".join(letter.contact))
    para(letter.date)
    para("\n".join(x for x in (letter.recipient, letter.company) if x))
    para(f"{letter.role} — {letter.company}", bold=True)
    para(letter.greeting)
    for p in letter.paragraphs:
        para(p)
    para(letter.closing)
    para(letter.name)
    buf = io.BytesIO()
    out.save(buf)
    return buf.getvalue()


def verify(pdf: bytes, letter: Letter) -> Report:
    """The CV's generic ATS checks, plus: one page, and every paragraph extracted whole."""
    report = verify_pdf(pdf, max_pages=1)
    flat = " ".join(report.text.split())
    lost = [p[:50] for p in letter.paragraphs if " ".join(p.split()) not in flat]
    report.checks.append(Check("paragraphs_extracted", not lost, "; ".join(lost)))
    return report


def file_stem(letter: Letter) -> str:
    safe = lambda s: re.sub(r"[^\w]+", "_", s, flags=re.UNICODE).strip("_")  # noqa: E731
    return f"{safe(letter.name)}_Cover_Letter_{safe(letter.company)}_{letter.lang.upper()}"


@dataclass
class BuiltLetter:
    path: Path
    format: str
    report: Report | None = None


def build(letter: Letter, out_dir: Path, formats: tuple[str, ...] = FORMATS) -> list[BuiltLetter]:
    out_dir.mkdir(parents=True, exist_ok=True)
    stem = file_stem(letter)
    built = []
    if "pdf" in formats:
        pdf = render_pdf(letter)
        path = out_dir / f"{stem}.pdf"
        path.write_bytes(pdf)
        built.append(BuiltLetter(path, "pdf", verify(pdf, letter)))
    if "docx" in formats:
        path = out_dir / f"{stem}.docx"
        path.write_bytes(render_docx(letter))
        built.append(BuiltLetter(path, "docx"))
    return built

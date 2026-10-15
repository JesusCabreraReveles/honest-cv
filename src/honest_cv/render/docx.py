"""DOCX rendering with python-docx: single column, plain paragraphs, no tables,
no headers or footers. Some job portals prefer Word files over PDF."""

from __future__ import annotations

import io

from docx import Document as DocxDocument
from docx.shared import Inches, Pt

from honest_cv.document import Document

FONT = "Calibri"  # metric-compatible with the Carlito used in the PDF


def render_docx(doc: Document) -> bytes:
    out = DocxDocument()
    for s in out.sections:
        s.top_margin = s.bottom_margin = Inches(0.45)
        s.left_margin = s.right_margin = Inches(0.55)
    normal = out.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(10)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.space_before = Pt(0)
    for name in ("Title", "Heading 1", "Heading 2", "List Bullet"):
        style = out.styles[name]
        style.font.name = FONT

    def para(text: str = "", *, bold: bool = False, size: float | None = None, style: str | None = None):
        p = out.add_paragraph(style=style)
        if text:
            run = p.add_run(text)
            run.bold = bold
            if size:
                run.font.size = Pt(size)
        return p

    def heading(text: str):
        h = out.add_heading(text, level=2)
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(2)

    def bullets(items: list[str]):
        for item in items:
            para(item, style="List Bullet")

    H = doc.headings
    title = out.add_heading(doc.name, level=1)
    title.paragraph_format.space_after = Pt(0)
    para(doc.title, bold=True, size=11)
    para(" · ".join(doc.contact))

    heading(H["summary"])
    para(doc.summary)
    if doc.competencies:
        heading(H["competencies"])
        para(" · ".join(doc.competencies))
    if doc.impact:
        heading(H["impact"])
        bullets(doc.impact)
    if doc.roles:
        heading(H["experience"])
        for r in doc.roles:
            p = para()
            p.paragraph_format.space_before = Pt(6)
            p.add_run(r.heading).bold = True
            p.add_run(f" · {r.dates}")
            if r.context:
                para(r.context)
            bullets(r.bullets)
    if doc.skills:
        heading(H["skills"])
        for s in doc.skills:
            p = para()
            p.add_run(f"{s.label}:").bold = True
            p.add_run(" " + ", ".join(s.skills))
    if doc.education:
        heading(H["education"])
        for e in doc.education:
            para(e)
    if doc.certifications:
        heading(H["certifications"])
        for c in doc.certifications:
            para(c)
    if doc.publications:
        heading(H["publications"])
        for x in doc.publications:
            para(x)
    if doc.courses:
        heading(H["courses"])
        for c in doc.courses:
            para(c)
    if doc.languages:
        heading(H["languages"])
        para(doc.languages)

    buf = io.BytesIO()
    out.save(buf)
    return buf.getvalue()

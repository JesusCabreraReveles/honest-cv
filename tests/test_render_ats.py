"""ATS invariants, checked on every build of the example profile."""

import copy
import io
import re

import pytest
from docx import Document as DocxDocument
from pdfminer.high_level import extract_text
from pypdf import PdfReader

from honest_cv import document
from honest_cv.build import build_cv
from honest_cv.render.docx import render_docx
from honest_cv.render.pdf import DENSITY, fit_pages, page_stats, render_pdf
from honest_cv.verify import verify_pdf

LIGATURES = re.compile("[ﬀ-ﬆ]")
CASES = [(v, l) for v in ("backend", "platform", "leadership") for l in ("en", "es")]


@pytest.fixture(scope="module")
def pdfs():
    from honest_cv.model import load
    from tests.conftest import ADA

    source = load(ADA)
    return {(v, l): (document.build(source, v, l), render_pdf(document.build(source, v, l))) for v, l in CASES}


@pytest.mark.parametrize("case", CASES)
def test_text_extractable_without_ligatures(pdfs, case):
    doc, pdf = pdfs[case]
    a = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages)
    b = extract_text(io.BytesIO(pdf))
    for text in (a, b):
        assert doc.name in text
        assert not LIGATURES.search(text)
        assert "­" not in text  # no soft hyphens
        for bullet in doc.impact + [b for r in doc.roles for b in r.bullets]:
            assert " ".join(bullet.split()) in " ".join(text.split())


@pytest.mark.parametrize("case", CASES)
def test_role_heading_keeps_its_dates_in_pdfminer(pdfs, case):
    doc, pdf = pdfs[case]
    lines = extract_text(io.BytesIO(pdf)).splitlines()
    for role in doc.roles:
        assert f"{role.heading} · {role.dates}" in lines


@pytest.mark.parametrize("case", CASES)
def test_reading_order_follows_source(pdfs, case):
    doc, pdf = pdfs[case]
    report = verify_pdf(pdf, doc=doc)
    assert next(c for c in report.checks if c.name == "reading_order").ok


def test_reading_order_catches_swapped_roles(pdfs):
    doc, pdf = pdfs[("backend", "en")]
    doc = copy.deepcopy(doc)
    doc.roles.reverse()
    report = verify_pdf(pdf, doc=doc)
    assert "reading_order" in {c.name for c in report.failures}


def test_only_bundled_fonts_embedded(pdfs):
    _, pdf = pdfs[("backend", "en")]
    reader = PdfReader(io.BytesIO(pdf))
    fonts = {str(f.get_object()["/BaseFont"]).split("+")[-1] for p in reader.pages for f in p["/Resources"]["/Font"].values()}
    assert fonts and all(f.startswith("Carlito") for f in fonts)


def test_no_images_in_pdf(pdfs):
    _, pdf = pdfs[("backend", "en")]
    assert all(not p.images for p in PdfReader(io.BytesIO(pdf)).pages)


def test_markup_in_user_text_is_literal(ada):
    doc = document.build(ada, "backend", "en")
    nasty = 'Costs $5 #set text(red) *bold* _x_ @ref <label> = heading // c'
    doc.summary = nasty
    text = extract_text(io.BytesIO(render_pdf(doc)))
    assert nasty in text


def test_fit_pages_steps_down_density(ada):
    doc = document.build(ada, "backend", "en")
    doc.max_pages = 1
    filler = doc.roles[0].bullets[0]
    # Grow the CV until the airiest density spills to a second page.
    for n in range(10, 120):
        doc.roles[0].bullets = [filler] * n
        airy, _ = page_stats(render_pdf(doc, DENSITY[0]))
        if airy > 1:
            break
    tight, _ = page_stats(render_pdf(doc, DENSITY[-1]))
    assert airy == 2 and tight == 1
    result = fit_pages(doc)
    assert result.level > 0 and result.pages == 1 and result.fits


def test_fit_pages_reports_overflow(ada):
    doc = document.build(ada, "backend", "en")
    doc.roles[0].bullets = [doc.roles[0].bullets[0]] * 150
    doc.max_pages = 1
    result = fit_pages(doc)
    assert not result.fits and result.pages > 1


@pytest.mark.parametrize(("lang", "heading"), [("en", "Courses"), ("es", "Cursos")])
def test_courses_get_their_own_section(ada, lang, heading):
    source = ada.model_copy(deep=True)
    source.profile.courses = [{"en": "Situational leadership course, 2021", "es": "Curso de liderazgo situacional, 2021"}]
    doc = document.build(source, "backend", lang)
    course = doc.courses[0]
    pdf = extract_text(io.BytesIO(render_pdf(doc)))
    docx = "\n".join(p.text for p in DocxDocument(io.BytesIO(render_docx(doc))).paragraphs)
    for text in (pdf, docx):
        assert text.index(doc.headings["certifications"]) < text.index(heading) < text.index(course)


def test_docx_is_single_column_without_tables_images_or_header(ada):
    for lang in ("en", "es"):
        d = DocxDocument(io.BytesIO(render_docx(document.build(ada, "backend", lang))))
        assert not d.tables
        assert not d.inline_shapes
        for s in d.sections:
            assert all(not p.text.strip() for p in s.header.paragraphs)
            assert all(not p.text.strip() for p in s.footer.paragraphs)


def test_build_cv_writes_files(ada, tmp_path):
    built = build_cv(ada, "backend", "es", tmp_path)
    assert [b.path.name for b in built] == ["Ada_Example_CV_backend_ES.pdf", "Ada_Example_CV_backend_ES.docx"]
    assert built[0].pages == 1 and built[0].fits


@pytest.mark.parametrize("term", ["CI/CD", "event-driven", "first-in-first-out"])
def test_compound_terms_never_split_across_lines(ada, term):
    doc = document.build(ada, "backend", "en")
    # Shift the term across every horizontal position so some copy lands at a line end.
    doc.summary = " ".join(f"{'x' * (i % 7 + 1)} {term}" for i in range(80))
    text = extract_text(io.BytesIO(render_pdf(doc)))
    assert text.count(term) >= 80  # every copy extracted whole
    for line in text.splitlines():
        assert not line.rstrip().endswith(("/", "-")), line


def test_punctuation_after_capitals_extracts_without_a_gap(ada):
    # Found with a real profile: kerning pulled the comma toward "P" and pypdf read
    # "TCP," as "TCP ,", so the bullet no longer matched. Kerning is off now.
    doc = document.build(ada, "backend", "en")
    sentence = "Moved from UDP to TCP, then to HTTP. Cut ERP, VPN and GPS costs; also Y, P and V."
    doc.roles[0].bullets[0] = sentence
    pdf = render_pdf(doc)
    pypdf_text = " ".join(" ".join(p.extract_text().split()) for p in PdfReader(io.BytesIO(pdf)).pages)
    assert sentence in pypdf_text
    assert verify_pdf(pdf, doc=doc).ok


def test_contact_items_never_split_across_lines(ada):
    doc = document.build(ada, "backend", "en")
    url = "https://a-very-long-personal-domain.example.systems/portfolio"
    # Push the URL across line ends with fillers of growing length.
    doc.contact = [f"{'x' * n}" for n in range(1, 25)] + [url]
    lines = [line.strip() for line in extract_text(io.BytesIO(render_pdf(doc))).splitlines()]
    assert any(url in line for line in lines)
